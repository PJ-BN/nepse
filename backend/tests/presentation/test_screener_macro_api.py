"""Screener, macro, breadth, mutual-funds and calendar routers over fake repositories."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from nepse.application.macro.ports import IssueRow
from nepse.application.mutual_funds.ports import FundSnapshot
from nepse.application.screener.filter_ast import OUTPUT_COLUMNS
from nepse.presentation.api.v1 import breadth, calendar, macro, mutual_funds, screener

ROW = {c: None for c in OUTPUT_COLUMNS} | {"symbol": "NABIL", "sector": "BANKING", "close": 500.0, "pbv": 1.2}


@pytest.fixture(autouse=True)
def no_db(override):
    override(screener.get_session, None)  # every router here reuses screener.get_session


class FakeScreenerRepo:
    queries = []

    def __init__(self, session):
        pass

    async def query(self, node, sort, limit):
        FakeScreenerRepo.queries.append((node, sort, limit))
        return [ROW]


@pytest.fixture
def screener_repo(monkeypatch):
    FakeScreenerRepo.queries = []
    monkeypatch.setattr(screener, "SqlScreenerRepository", FakeScreenerRepo)
    return FakeScreenerRepo


async def test_run_json(client, screener_repo):
    body = {"filter": {"field": "pbv", "op": "<", "value": 2}, "sort": ["-pbv"], "limit": 5}
    r = (await client.post("/api/v1/screener/run", json=body)).json()
    assert r["columns"] == list(OUTPUT_COLUMNS) and r["count"] == 1 and r["rows"][0]["symbol"] == "NABIL"
    [(node, sort, limit)] = screener_repo.queries
    assert node is not None and limit == 5 and len(sort) == 1


async def test_run_preset(client, screener_repo):
    assert (await client.post("/api/v1/screener/run", json={"preset": "undervalued-banks"})).status_code == 200
    assert (await client.post("/api/v1/screener/run", json={"preset": "nope"})).status_code == 404


@pytest.mark.parametrize("body", [
    {"filter": {"field": "bogus", "op": "<", "value": 1}},
    {"filter": {"field": "pbv", "op": "; DROP TABLE x", "value": 1}},
    {"sort": ["not_a_column"]},
])
async def test_run_rejects_invalid_filters(client, screener_repo, body):
    assert (await client.post("/api/v1/screener/run", json=body)).status_code == 422
    assert screener_repo.queries == []  # rejected before the repository is touched


async def test_run_limit_bounds(client, screener_repo):
    for limit in (0, 1001):
        assert (await client.post("/api/v1/screener/run", json={"limit": limit})).status_code == 422


async def test_run_csv(client, screener_repo):
    r = await client.post("/api/v1/screener/run", json={"format": "csv"})
    assert r.headers["content-type"].startswith("text/csv")
    assert 'filename="screener.csv"' in r.headers["content-disposition"]
    header, line = r.text.splitlines()
    assert header.split(",") == list(OUTPUT_COLUMNS) and line.startswith("NABIL,BANKING,")


async def test_run_xlsx(client, screener_repo):
    r = await client.post("/api/v1/screener/run", json={"format": "xlsx"})
    assert r.headers["content-type"] == screener.XLSX
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == list(OUTPUT_COLUMNS) and ws["A2"].value == "NABIL"


async def test_presets(client):
    presets = (await client.get("/api/v1/screener/presets")).json()
    assert {"undervalued-banks", "rsi-bullish-divergence"} <= {p["id"] for p in presets}


class FakeMacroRepo:
    def __init__(self, session):
        pass

    async def series(self, indicators, since):
        return {"cd_ratio": [(date(2025, 1, 1), 80.0)]}

    async def index_closes(self, codes, since):
        return {"NEPSE": [(date(2025, 1, 1), 2000.0)]}

    async def issues(self):
        return [IssueRow("ABC Hydro", "IPO", "PIPELINE", sector="HYDRO", units=1e6)]


async def test_macro(client, monkeypatch):
    monkeypatch.setattr(macro, "SqlMacroRepository", FakeMacroRepo)
    radar = (await client.get("/api/v1/macro/radar", params={"days": 90})).json()
    assert radar["nepse"] == [{"date": "2025-01-01", "value": 2000.0}]
    assert radar["series"]["cd_ratio"][0]["value"] == 80.0 and radar["regime"] is None  # too little history
    rrg = (await client.get("/api/v1/macro/rrg", params={"timeframe": "weekly", "tail": 3})).json()
    assert rrg["benchmark"] == "NEPSE" and rrg["timeframe"] == "weekly" and isinstance(rrg["sectors"], list)
    [issue] = (await client.get("/api/v1/macro/ipo-pipeline")).json()
    assert issue["company"] == "ABC Hydro" and issue["stage"] == "PIPELINE"


async def test_macro_validation(client, monkeypatch):
    monkeypatch.setattr(macro, "SqlMacroRepository", FakeMacroRepo)
    for url, params in [("radar", {"days": 29}), ("radar", {"days": 3651}), ("rrg", {"timeframe": "hourly"}),
                        ("rrg", {"tail": 0}), ("rrg", {"tail": 31})]:
        assert (await client.get(f"/api/v1/macro/{url}", params=params)).status_code == 422, (url, params)


class FakeBreadthRepo:
    closes: dict = {}

    def __init__(self, session):
        pass

    async def adjusted_closes(self, days):
        return FakeBreadthRepo.closes


async def test_breadth(client, monkeypatch):
    monkeypatch.setattr(breadth, "SqlScreenerRepository", FakeBreadthRepo)
    FakeBreadthRepo.closes = {}
    assert (await client.get("/api/v1/breadth")).status_code == 404
    d = [date(2025, 1, i) for i in (1, 2, 3)]
    FakeBreadthRepo.closes = {"UP": list(zip(d, [10.0, 11.0, 12.0])), "DOWN": list(zip(d, [10.0, 9.0, 8.0])),
                              "FLAT": list(zip(d, [10.0, 10.0, 10.0]))}
    body = (await client.get("/api/v1/breadth")).json()
    assert (body["advances"], body["declines"], body["unchanged"]) == (1, 1, 1)
    assert body["as_of"] == "2025-01-03" and len(body["series"]) == 2 and set(body["pct_above"]) == {"sma_20", "sma_50", "sma_200"}


async def test_mutual_funds(client, monkeypatch):
    class Repo:
        def __init__(self, session):
            pass

        async def snapshots(self):
            return [FundSnapshot("NOFUND", "No NAV Fund", None, None, None, None),
                    FundSnapshot("DEEP", "Deep Discount", 7.0, 10.0, date(2025, 1, 1), date.today().replace(year=date.today().year + 3)),
                    FundSnapshot("SOON", "Soon Maturing", 8.0, 10.0, date(2025, 1, 1), date.today())]

    monkeypatch.setattr(mutual_funds, "SqlMutualFundRepository", Repo)
    rows = (await client.get("/api/v1/mutual-funds")).json()
    assert [r["symbol"] for r in rows] == ["SOON", "DEEP", "NOFUND"]  # flagged first, no-NAV last
    assert rows[0]["flagged"] and rows[0]["nav_discount_pct"] == -20.0 and not rows[1]["flagged"]
    assert rows[2]["nav_discount_pct"] is None and rows[2]["days_to_maturity"] is None


async def test_calendar_validation_and_passthrough(client, monkeypatch):
    seen = {}

    class Repo:
        def __init__(self, session):
            pass

        async def announcements(self, limit):
            seen["limit"] = limit
            return []

    monkeypatch.setattr(calendar, "SqlCalendarRepository", Repo)
    assert (await client.get("/api/v1/calendar/announcements", params={"limit": 7})).json() == []
    assert seen["limit"] == 7
    for path, params in [("announcements", {"limit": 0}), ("announcements", {"limit": 201}),
                         ("dividends", {"limit": 501}), ("rights", {"limit": 0}),
                         ("book-closures", {"days": 0}), ("book-closures", {"days": 366})]:
        assert (await client.get(f"/api/v1/calendar/{path}", params=params)).status_code == 422, (path, params)
