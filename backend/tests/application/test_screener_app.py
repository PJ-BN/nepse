import pytest

from nepse.application.screener.filter_ast import (
    FIELDS, OUTPUT_COLUMNS, Condition, FieldRef, FilterError, Group, SortKey, parse_filter, parse_limit, parse_sort,
)
from nepse.application.screener.run_screener import PRESETS, run_preset, run_screener, to_csv, to_xlsx


def test_parse_nested_tree():
    node = parse_filter({"and": [
        {"field": "close", "op": ">", "value": {"field": "sma_200"}},
        {"or": [
            {"field": "sector", "op": "in", "value": ["HYDRO", "BANKING"]},
            {"field": "rsi_14", "op": "between", "value": [20, 30.5]},
        ]},
    ]})
    assert node == Group("and", (
        Condition("close", ">", FieldRef("sma_200")),
        Group("or", (
            Condition("sector", "in", ("HYDRO", "BANKING")),
            Condition("rsi_14", "between", (20.0, 30.5)),
        )),
    ))
    assert parse_filter(None) is None and parse_filter({}) is None


@pytest.mark.parametrize("bad", [
    # unknown / injected field names
    {"field": "close; DROP TABLE companies", "op": ">", "value": 1},
    {"field": "pg_sleep(10)", "op": "=", "value": 1},
    {"field": "close", "op": ">", "value": {"field": "1=1 OR close"}},
    {"field": ["close"], "op": ">", "value": 1},
    # unknown / injected ops
    {"field": "close", "op": "> 0 OR 1=1 --", "value": 1},
    {"field": "close", "op": "like", "value": 1},
    {"field": "sector", "op": "<", "value": "A"},          # ordering op on text
    {"field": "close", "op": None, "value": 1},
    # wrongly typed values
    {"field": "close", "op": ">", "value": "100"},
    {"field": "close", "op": ">", "value": "1; DROP TABLE companies"},
    {"field": "close", "op": ">", "value": True},
    {"field": "close", "op": ">", "value": float("nan")},
    {"field": "close", "op": ">", "value": None},
    {"field": "sector", "op": "=", "value": 5},
    {"field": "sector", "op": "=", "value": "x" * 65},
    {"field": "close", "op": "between", "value": [1]},
    {"field": "close", "op": "between", "value": [1, "2"]},
    {"field": "close", "op": "in", "value": []},
    {"field": "close", "op": "in", "value": list(range(201))},
    {"field": "close", "op": "in", "value": 5},
    {"field": "close", "op": ">", "value": {"field": "sector"}},  # num vs text
    {"field": "close", "op": ">", "value": {"field": "sma_200", "raw": "x"}},
    # structural
    {"field": "close", "op": ">", "value": 1, "extra": "x"},
    {"and": []},
    {"and": {"field": "close"}},
    {"xor": [{"field": "close", "op": ">", "value": 1}]},
    {"and": [{"field": "close", "op": ">", "value": 1}], "or": []},
    ["close"],
    "close > 1",
])
def test_rejects_invalid_and_injection_attempts(bad):
    with pytest.raises(FilterError):
        parse_filter(bad)


def test_rejects_huge_or_deep_trees():
    leaf = {"field": "close", "op": ">", "value": 1}
    with pytest.raises(FilterError):
        parse_filter({"or": [leaf] * 60})
    deep: dict = leaf
    for _ in range(10):
        deep = {"and": [deep]}
    with pytest.raises(FilterError):
        parse_filter(deep)


def test_sort_and_limit():
    assert parse_sort(None) == [SortKey("symbol", False)]
    assert parse_sort(["-rsi_14", "symbol"]) == [SortKey("rsi_14", True), SortKey("symbol", False)]
    assert parse_sort("close") == [SortKey("close", False)]
    for bad in (["close desc"], ["--close"], ["symbol; DROP"], [1], ["a", "b", "c", "d"], {"f": 1}):
        with pytest.raises(FilterError):
            parse_sort(bad)
    assert parse_limit(5) == 5
    for bad in (0, 1001, "10", True, 1.5):
        with pytest.raises(FilterError):
            parse_limit(bad)


def test_output_columns_cover_every_field():
    assert set(FIELDS) <= set(OUTPUT_COLUMNS)


class FakeRepo:
    def __init__(self, rows):
        self.rows, self.calls = rows, []

    async def query(self, node, sort, limit):
        self.calls.append((node, sort, limit))
        return self.rows[:limit]

    async def refresh(self):
        pass


async def test_run_screener_validates_before_querying():
    repo = FakeRepo([{"symbol": "NABIL"}])
    assert await run_screener(repo, {"field": "close", "op": ">", "value": 1}, ["-close"], 10) == [{"symbol": "NABIL"}]
    with pytest.raises(FilterError):
        await run_screener(repo, {"field": "nope", "op": ">", "value": 1})
    assert len(repo.calls) == 1


@pytest.mark.parametrize("preset_id", list(PRESETS))
async def test_presets_are_valid(preset_id):
    repo = FakeRepo([])
    await run_preset(repo, preset_id)
    node, sort, limit = repo.calls[0]
    assert node is not None and sort and limit == 100


def test_csv_export():
    out = to_csv([{"symbol": "NABIL", "sector": "BANKING", "close": 512.5, "junk": "ignored"}])
    header, row = out.splitlines()
    assert header.split(",") == list(OUTPUT_COLUMNS)
    assert row.startswith("NABIL,BANKING,,512.5")


def test_xlsx_export():
    from io import BytesIO

    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(to_xlsx([{"symbol": "NABIL", "close": 512.5, "junk": 1}]))).active
    assert [c.value for c in ws[1]] == list(OUTPUT_COLUMNS)
    assert ws[2][0].value == "NABIL" and ws[2][list(OUTPUT_COLUMNS).index("close")].value == 512.5
